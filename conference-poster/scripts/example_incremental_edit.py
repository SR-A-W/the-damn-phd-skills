# EXAMPLE ONLY — taken from the PLE poster for COLM 2026. Paths, numbers and panel names are paper-specific;
# copy the structure (LAYOUT dict + helpers + fixed type scale), not the content.
"""v10 = v9 + (1) /think navy & /no_think green in all body text, (2) fuller result conclusions, (3) Demystifying QR bottom-right."""
import re, copy
from pptx import Presentation
from pptx.util import Pt, Inches, Emu
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from lxml import etree

SRC, OUT = "PLE_poster_v9.pptx", "PLE_poster_v10.pptx"
A = "/tmp/claude-1097325/-scratch-pioneer-users-sxw992-hybrid-expert-thinking/ee9492a4-43e9-48da-a3fd-87d8aa1ffa64/scratchpad/poster_assets"
NAVY = RGBColor(0x1F, 0x3A, 0x5F); GREEN = RGBColor(0x2E, 0x7D, 0x32); INK = RGBColor(0x22, 0x22, 0x22)
prs = Presentation(SRC); s = prs.slides[0]
TOK = re.compile(r"(/no_think|/think)")

def colorize_paragraph(p):
    """Split runs so that /think and /no_think get their own colored runs; other formatting is inherited."""
    for r in list(p.runs):
        t = r.text
        if not TOK.search(t): continue
        parts = [x for x in TOK.split(t) if x != ""]
        prev = r._r
        for i, part in enumerate(parts):
            if i == 0:
                r.text = part; cur_r = r
            else:
                new_r = copy.deepcopy(r._r); prev.addnext(new_r); prev = new_r
                from pptx.text.text import _Run
                cur_r = _Run(new_r, p); cur_r.text = part
            if part == "/think": cur_r.font.color.rgb = NAVY
            elif part == "/no_think": cur_r.font.color.rgb = GREEN
            elif i > 0:  # restore default color for the tail piece (copied run may carry a token color)
                cur_r.font.color.rgb = INK if (r.font.color and r.font.color.type is not None and r.font.color.rgb in (NAVY, GREEN)) else cur_r.font.color.rgb if cur_r.font.color.type is not None else INK
            if i == 0 and part not in ("/think", "/no_think"):
                pass
        # first piece: if it was a token it is colored above; else leave as is

# (1) all text frames incl. tables; skip 44pt headings
for sh in s.shapes:
    frames = []
    if sh.has_text_frame: frames.append(sh.text_frame)
    if sh.has_table:
        for row in sh.table.rows:
            for c in row.cells: frames.append(c.text_frame)
    for tf in frames:
        for p in tf.paragraphs:
            if any(r.font.size and r.font.size.pt >= 40 for r in p.runs): continue
            colorize_paragraph(p)

# (2) fuller conclusions; shrink Fig.3 slightly to make room
by_name = {sh.name: sh for sh in s.shapes}
fig = by_name["Picture 29"]; ratio = fig.width / fig.height
fig.height = Inches(6.1); fig.width = int(fig.height * ratio); fig.left = Inches(1.0 + (21.0 - fig.width / 914400) / 2)
cap = by_name["TextBox 30"]; cap.top = Inches(17.6 + 6.1)
def set_bullet(p):
    pPr = p._p.get_or_add_pPr()
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)): pPr.remove(el)
    pPr.set("marL", str(int(Inches(0.45)))); pPr.set("indent", str(-int(Inches(0.45))))
    etree.SubElement(pPr, qn("a:buChar")).set("char", "•")
lines = [
 ("Maintains /think performance.", " With both modes trained on the same data, PLE reaches 61.3% on AIME24 in /think mode versus 60.0% for the SFT-only recipe, so giving each mode its own feed-forward path costs nothing on hard reasoning."),
 ("Improves /no_think performance.", " /no_think accuracy rises from 35.3% (SFT-only) to 44.7%, well above the off-the-shelf hybrid Qwen3-4B (23.3%): a dedicated no-think expert makes direct answering stronger, not weaker."),
 ("Separates the modes.", " SFT-only halves /no_think length (19,799 → 8,665 tokens) but leaves reflective content almost untouched (6.91 → 6.01); PLE halves the length again (4,101) and removes the reflective tokens (0.35) — both leakage signals drop at once."),
]
for sh in s.shapes:
    if sh.has_text_frame and "Maintains /think performance" in sh.text_frame.text:
        tf = sh.text_frame
        for p in list(tf.paragraphs)[1:]: p._p.getparent().remove(p._p)
        p0 = tf.paragraphs[0]
        for r in list(p0.runs): r._r.getparent().remove(r._r)
        for i, (label, body) in enumerate(lines):
            p = p0 if i == 0 else tf.add_paragraph(); p.space_after = Pt(4); set_bullet(p)
            for txt, b in ((label, True), (body, False)):
                r = p.add_run(); r.text = txt; r.font.name = "Calibri"; r.font.size = Pt(24); r.font.bold = b; r.font.color.rgb = INK
            colorize_paragraph(p)
        sh.top = Inches(24.35); sh.height = Inches(3.45)
# (3) Demystifying QR in the bottom-right cluster; COLM logo shifts right and shrinks a bit
colm_bottom = [sh for sh in s.shapes if sh.shape_type == 13 and Emu(sh.top).inches > 25 and Emu(sh.left).inches > 48][0]
r = colm_bottom.width / colm_bottom.height; colm_bottom.height = Inches(2.0); colm_bottom.width = int(Inches(2.0) * r)
colm_bottom.left = Inches(54.7 - colm_bottom.width / 914400); colm_bottom.top = Inches(25.45)
s.shapes.add_picture(f"{A}/qr_demystifying.png", Inches(48.7), Inches(25.4), Inches(2.0), Inches(2.0))
tb = s.shapes.add_textbox(Inches(48.3), Inches(27.4), Inches(2.8), Inches(0.45)); tf = tb.text_frame; tf.margin_left = tf.margin_right = 0
p = tf.paragraphs[0]; p.alignment = 2; rr = p.add_run(); rr.text = "Prior work"; rr.font.name = "Calibri"; rr.font.size = Pt(18); rr.font.bold = True; rr.font.color.rgb = NAVY
prs.save(OUT); print("saved", OUT)
