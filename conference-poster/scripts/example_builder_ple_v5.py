# EXAMPLE ONLY — taken from the PLE poster for COLM 2026. Paths, numbers and panel names are paper-specific;
# copy the structure (LAYOUT dict + helpers + fixed type scale), not the content.
"""PLE COLM 2026 poster v5 — consistent type styles (BODY vs CAPTION), bordered panels,
white panels for figures, uniform logo heights, Ablations as text-left / figures-right."""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = os.path.dirname(os.path.abspath(__file__)); POSTER = os.path.dirname(HERE)
A = "/tmp/claude-1097325/-scratch-pioneer-users-sxw992-hybrid-expert-thinking/ee9492a4-43e9-48da-a3fd-87d8aa1ffa64/scratchpad/poster_assets"
LOGOS = f"{POSTER}/logos"
TEMPLATE = f"{POSTER}/Main_Conference_template_COLM.pptx"; OUT = f"{POSTER}/PLE_poster_v5.pptx"

NAVY = RGBColor(0x1F, 0x3A, 0x5F); GREEN = RGBColor(0x2E, 0x7D, 0x32); RED = RGBColor(0xB0, 0x1E, 0x1E)
INK = RGBColor(0x22, 0x22, 0x22); GRAY = RGBColor(0x55, 0x55, 0x55); CAPGRAY = RGBColor(0x60, 0x60, 0x60)
PANEL = RGBColor(0xF3, 0xF5, 0xF8); WHITE = RGBColor(0xFF, 0xFF, 0xFF); BORDER = RGBColor(0x9A, 0xA8, 0xB8)
FONT = "Calibri"
# ---- type scale (poster) ----
H_SIZE, BODY, BODY_S, CAP, TBL = 44, 28, 26, 20, 22

M = 1.0; GAP = 0.6
ROW1_Y, ROW1_H = 4.9, 10.9
ROW2_Y, ROW2_H = 16.3, 11.6
L = {
    "title":   (11.0, 0.4, 34.0, 2.3), "authors": (9.0, 2.75, 38.0, 0.9), "affil": (9.0, 3.6, 38.0, 0.8),
    "p_abs":   (M, ROW1_Y, 14.4, ROW1_H),
    "p_arch":  (M + 15.0, ROW1_Y, 24.0, ROW1_H),
    "p_mot":   (M + 39.6, ROW1_Y, 14.4, ROW1_H),
    "p_res":   (M, ROW2_Y, 21.0, ROW2_H),
    "p_abl":   (M + 21.6, ROW2_Y, 17.6, ROW2_H),
    "p_gen":   (M + 39.8, ROW2_Y, 14.2, ROW2_H),
}

prs = Presentation(TEMPLATE); slide = prs.slides[0]
for sh in list(slide.shapes): sh._element.getparent().remove(sh._element)

def box(x, y, w, h, fill=PANEL, line=BORDER, lw=3.0):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = 0.02; s.shadow.inherit = False
    if fill is None: s.fill.background()
    else: s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None: s.line.fill.background()
    else: s.line.color.rgb = line; s.line.width = Pt(lw)
    return s

def text(x, y, w, h, runs, size=BODY, bold=False, italic=False, color=INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, para_space=6):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.12); tf.margin_top = tf.margin_bottom = Inches(0.06)
    for i, p in enumerate(runs if isinstance(runs, list) else [runs]):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); para.alignment = align; para.space_after = Pt(para_space)
        for seg in (p if isinstance(p, list) else [(p, {})]):
            t, o = seg if isinstance(seg, tuple) else (seg, {})
            r = para.add_run(); r.text = t; r.font.name = FONT
            r.font.size = Pt(o.get("size", size)); r.font.bold = o.get("bold", bold); r.font.italic = o.get("italic", italic)
            r.font.color.rgb = o.get("color", color)
    return tb

def caption(x, y, w, h, s):  # figure/table caption style — distinct from body
    return text(x, y, w, h, s, size=CAP, italic=True, color=CAPGRAY)

def pic(path, x, y, w=None, h=None):
    return slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)

def panel(key, heading, figure=False):
    x, y, w, h = L[key]; box(x, y, w, h, fill=WHITE if figure else PANEL)
    text(x + 0.25, y + 0.12, w - 0.5, 1.05, heading, size=H_SIZE, bold=True, color=NAVY)
    return x, y, w, h

def table(x, y, w, h, rows, col_w=None, size=TBL, bold_last=False):
    nr, nc = len(rows), len(rows[0])
    gt = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(h)).table
    if col_w:
        for j, cw in enumerate(col_w): gt.columns[j].width = Inches(cw)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            c = gt.cell(i, j); c.text = ""; p = c.text_frame.paragraphs[0]; r = p.add_run(); r.text = str(val)
            r.font.name = FONT; r.font.size = Pt(size); p.alignment = PP_ALIGN.CENTER if j else PP_ALIGN.LEFT
            c.margin_top = c.margin_bottom = Inches(0.05); c.fill.solid()
            if i == 0: c.fill.fore_color.rgb = NAVY; r.font.bold = True; r.font.color.rgb = WHITE
            else:
                c.fill.fore_color.rgb = WHITE if i % 2 else RGBColor(0xEA, 0xEE, 0xF4); r.font.color.rgb = INK
                if bold_last and i == nr - 1: r.font.bold = True; r.font.color.rgb = NAVY
    return gt

# ================= header =================
x, y, w, h = L["title"]
text(x, y, w, h, "Path-Lock Expert: Separating Reasoning Mode in Hybrid Thinking via Architecture-Level Separation",
     size=62, bold=True, color=NAVY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(*L["authors"], "Shouren Wang¹*, Wang Yang¹*, Chuang Ma², Debargha Ganguly¹, Vikash Singh¹, Chaoda Song¹, Xinpeng Li¹, Xianxuan Long³, Vipin Chaudhary¹†, Xiaotian Han¹†",
     size=28, color=INK, align=PP_ALIGN.CENTER)
text(*L["affil"], "¹Case Western Reserve University   ²Kyoto University, NII LLMC   ³Michigan State University      *Equal contribution   †Corresponding authors",
     size=22, color=GRAY, align=PP_ALIGN.CENTER)
# logos at a uniform 2.0" height (TAPP AI is a wide wordmark -> 1.0" tall, placed below)
LH = 2.0
pic(f"{LOGOS}/cwru.png", 1.0, 0.9, h=LH)                       # 560x208  -> w 5.4
pic(f"{LOGOS}/kyoto_university.png", 6.8, 0.9, h=LH)           # 899x506  -> w 3.6
pic(f"{LOGOS}/nii.png", 46.6, 0.9, h=LH)                       # 1600x1039 -> w 3.1
box(50.1, 0.9, 4.6, LH, fill=WHITE, line=RGBColor(0xCC, 0xCC, 0xCC), lw=1.5)
text(50.1, 0.9, 4.6, LH, "MSU logo", size=18, color=RGBColor(0x99, 0x99, 0x99), align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
pic(f"{LOGOS}/tapp_ai.png", 47.6, 3.15, h=1.05)                 # 700x150 -> w 4.9

# ================= row 1 =================
x, y, w, h = panel("p_abs", "Abstract")
text(x + 0.25, y + 1.25, w - 0.5, 3.2,
     [[("Hybrid-thinking LLMs expose ", {}), ("/think", {"bold": True, "color": NAVY}), (" and ", {}), ("/no_think", {"bold": True, "color": GREEN}),
       (" modes but do not separate them cleanly: even in /no_think they emit long, self-reflective answers — ", {}),
       ("reasoning leakage", {"bold": True, "color": NAVY}), (". Better data and multi-stage SFT reduce it, yet it remains because both modes share one set of feed-forward weights. ", {}),
       ("Path-Lock Expert (PLE)", {"bold": True, "color": NAVY}),
       (" gives each layer two expert MLPs — one per mode — behind shared attention; a deterministic control-token router locks one path for the whole sequence.", {})]],
     size=BODY_S)
iw = w - 0.9; ih = iw / 4.25; fy = y + 4.55
pic(f"{A}/motivation.png", x + 0.45, fy, iw, ih)
caption(x + 0.25, fy + ih + 0.02, w - 0.5, 0.9, "Figure 1. Qwen3-8B on an AIME24 problem: the /no_think answer still says “Wait” and is wrong, while /think solves it.")
by = fy + ih + 0.95
text(x + 0.25, by, w - 0.5, h - (by - y) - 0.15,
     [[("▸ ", {"color": NAVY}), ("Router-free dual-expert MLP", {"bold": True, "color": NAVY}), (" at the dense model’s per-token compute.", {})],
      [("▸ ", {"color": NAVY}), ("Stronger /no_think:", {"bold": True, "color": NAVY}), (" 17× fewer reflective tokens, 2× shorter, +9.3 pts; /think preserved.", {})],
      [("▸ ", {"color": NAVY}), ("Evidence", {"bold": True, "color": NAVY}), (" on 4 backbones, 2 corpora, Llama-3.1-8B, full MMLU, LLM-judge check.", {})]],
     size=24, para_space=4)

x, y, w, h = panel("p_arch", "Path-Lock Expert (PLE)", figure=True)
iw = w - 0.6; ih = iw / 2.69
pic(f"{A}/ple_architecture.png", x + 0.3, y + 1.3, iw, ih)
caption(x + 0.3, y + 1.3 + ih + 0.1, iw, h - (1.3 + ih + 0.2),
        "Figure 2. Two expert MLPs per decoder layer (think / no_think) behind shared attention, embeddings, normalization and LM head. A one-time deterministic routing decision from the control token locks one path for the whole sequence — no learned router, no routing loss, same per-token compute as the dense model.")

x, y, w, h = panel("p_mot", "Motivation: Reasoning Leakage")
text(x + 0.25, y + 1.25, w - 0.5, 3.4,
     [[("Asked ", {}), ("not", {"bold": True}), (" to think, released hybrid models still emit reflective markers (", {}),
       ("Wait…", {"bold": True, "color": RED}), (", ", {}), ("Hmm…", {"bold": True, "color": RED}),
       (") and long outputs: their /no_think behaves like a softened /think, not a genuine direct-answer mode. Training-level fixes reduce this but never remove it.", {})]], size=BODY_S)
rows = [["", "Qwen2.5-7B", "", "Qwen3-4B", "", "Qwen3-8B", ""],
        ["/no_think", "Len.", "#Refl.", "Len.", "#Refl.", "Len.", "#Refl."],
        ["MATH500", "703", "0.00", "1,583", "0.04", "1,296", "0.01"],
        ["AIME24", "1,729", "0.00", "12,506", "0.17", "12,109", "0.07"],
        ["GPQA", "775", "0.00", "2,679", "0.05", "2,943", "1.85"]]
ty = y + 4.3
gt = table(x + 0.25, ty, w - 0.5, 4.2, rows, col_w=[2.5, 2.0, 1.9, 2.0, 1.9, 1.9, 1.7], size=24)
for j0 in (1, 3, 5): gt.cell(0, j0).merge(gt.cell(0, j0 + 1))
for j in range(7):
    c = gt.cell(1, j); c.fill.solid(); c.fill.fore_color.rgb = RGBColor(0xDD, 0xE4, 0xEE)
    for r in c.text_frame.paragraphs[0].runs: r.font.bold = True; r.font.color.rgb = NAVY
caption(x + 0.25, ty + 4.25, w - 0.5, 0.8, "Table 1. /no_think outputs: a pure instruct model (Qwen2.5-7B-Instruct) vs. off-the-shelf hybrid models.")
text(x + 0.25, ty + 5.1, w - 0.5, h - (ty + 5.1 - y) - 0.1,
     "Hybrid /no_think answers are up to 7× longer than a genuine direct-answer mode and still contain reflective tokens.", size=BODY_S, bold=True, color=NAVY)

# ================= row 2 =================
x, y, w, h = panel("p_res", "Main results (AIME24, Qwen3-4B backbone)", figure=True)
ratio = 4830 / (3156 * 0.57); ih = 6.9; iw = ih * ratio; ox = x + (w - iw) / 2
pic(f"{A}/exp_top.png", ox, y + 1.25, iw, ih)
caption(x + 0.25, y + 1.25 + ih + 0.02, w - 0.5, 0.6, "Figure 3. Accuracy, output length and reflective tokens per answer; ★ marks the best value in each group.")
text(x + 0.25, y + 1.25 + ih + 0.75, w - 0.5, h - (1.25 + ih + 0.85),
     [[("vs. SFT-only: ", {"bold": True, "color": NAVY}), ("17× fewer reflective tokens in /no_think (6.01 → 0.35), 2× shorter outputs (8,665 → 4,101 tokens), +9.3 pts /no_think accuracy (35.3% → 44.7%), /think accuracy preserved (61.3% vs. 60.0%). The same picture holds on MATH500, GPQA-Diamond, MMLU-STEM and on the Qwen2.5-7B backbone.", {})]],
     size=BODY_S)

x, y, w, h = panel("p_abl", "Ablations & Takeaways", figure=True)
tw = 5.2; fx = x + 0.25 + tw + 0.3; fw = w - tw - 0.8
ih1 = fw / 3.23; ih2 = fw / 3.06
f1y = y + 1.25; f2y = f1y + ih1 + 0.75
pic(f"{A}/abl_base_composite_aime24.png", fx, f1y, fw, ih1)
caption(fx, f1y + ih1, fw, 0.65, "Figure 4. Backbone initialization (AIME24).")
pic(f"{A}/abl_dataset_top.png", fx, f2y, fw, ih2)
caption(fx, f2y + ih2, fw, 0.65, "Figure 5. Training data: Superior-reasoning vs. OpenR1 (AIME24, Qwen3-4B).")
text(x + 0.25, y + 1.3, tw, h - 1.5,
     [[("Backbone choice matters. ", {"bold": True, "color": NAVY}), ("Residual leakage orders with the backbone’s own reasoning lineage; a raw pretrained base collapses on hard tasks.", {})],
      [("Data difficulty must match capacity. ", {"bold": True, "color": NAVY}), ("The harder Superior-reasoning corpus wins in both modes on Qwen3-4B; OpenR1 stays competitive on Qwen2.5-7B.", {})],
      [("Separation is architectural. ", {"bold": True, "color": NAVY}), ("Separate feed-forward paths isolate mode-specific behavior while attention stays shared.", {})]],
     size=BODY_S, para_space=14)

x, y, w, h = panel("p_gen", "Generalization")
text(x + 0.25, y + 1.25, w - 0.5, 0.6, "Llama-3.1-8B, /no_think accuracy (%)", size=TBL, bold=True, color=NAVY)
rows = [["Model", "MATH500", "AIME24", "GPQA", "MMLU-STEM"], ["Instruct", "49.3", "6.0", "30.5", "53.0"],
        ["SFT-only", "47.2", "4.3", "26.7", "53.4"], ["PLE (ours)", "56.0", "12.0", "32.0", "76.7"]]
table(x + 0.25, y + 1.9, w - 0.5, 2.7, rows, col_w=[3.3, 2.5, 2.3, 2.3, 3.3], bold_last=True)
text(x + 0.25, y + 4.8, w - 0.5, 0.6, "Full MMLU (14,042 q.), Qwen3-4B, /no_think", size=TBL, bold=True, color=NAVY)
rows = [["Model", "Acc. (%)", "Len.", "#Refl./Ans."], ["Instruct", "79.79", "265", "0.02"],
        ["SFT-only", "75.58", "189", "0.16"], ["PLE (ours)", "75.60", "117", "0.00"]]
table(x + 0.25, y + 5.45, w - 0.5, 2.7, rows, col_w=[4.0, 3.3, 3.0, 3.4], bold_last=True)
text(x + 0.25, y + 8.3, w - 0.5, 0.9, "LLM judge (GPT-OSS-20B) agrees with the reflective-token metric in 92–100% of cases.", size=22, color=INK)
by = y + h - 2.5
pic(f"{LOGOS}/ple_logo_draft1.png", x + 0.25, by - 0.1, h=2.4)
pic(f"{A}/qr_arxiv.png", x + 3.3, by, 2.0, 2.0); text(x + 3.3, by + 2.0, 2.0, 0.4, "arXiv", size=18, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
pic(f"{A}/qr_code.png", x + 5.6, by, 2.0, 2.0);  text(x + 5.6, by + 2.0, 2.0, 0.4, "Code", size=18, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
box(x + 8.0, by + 0.1, w - 8.25, 2.0, fill=NAVY, line=None)
text(x + 8.0, by + 0.1, w - 8.25, 2.0, [[("COLM 2026", {"size": 28, "bold": True, "color": WHITE})], [("San Francisco", {"size": 17, "color": WHITE})]],
     align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

prs.save(OUT); print("saved", OUT)
